from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else Path(__file__).resolve().parent
src=ROOT/'Escenario-creatividad-validacion.md'
doc=Document()
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.75); s.bottom_margin=Inches(.7)
s.left_margin=Inches(.8); s.right_margin=Inches(.8)
s.header_distance=Inches(.3); s.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]
    st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
title=doc.styles['Title']; title.font.size=Pt(25); title.font.bold=True
title.paragraph_format.space_after=Pt(12)
doc.styles['Subtitle'].font.size=Pt(13)
for n,sz in [('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
    st=doc.styles[n]; st.font.size=Pt(sz); st.font.bold=True
    st.paragraph_format.space_before=Pt(14); st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=False
header=s.header.paragraphs[0]
header.text=next(line[2:] for line in src.read_text().splitlines() if line.startswith('# '))
header.runs[0].font.size=Pt(8)
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('00G-R01 · Non-canonical research  ·  '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title=header.text
doc.core_properties.subject='Non-canonical working scenario and Ecosystem Awareness candidacy'
doc.core_properties.author='Iván Abril Palma'
doc.core_properties.keywords='creativity, validation, agents, Ecosystem Awareness, Napoleon, Hugging Face'

def hyperlink(p,label,url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'1F4E79'); pr.append(color)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u)
    r.append(pr); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

def inline(p,text):
    pat=r'(\*\*[^*]+\*\*|\[[^\]]+\]\(https?://[^)]+\))'
    for token in re.split(pat,text):
        if not token: continue
        if token.startswith('**'):
            p.add_run(token[2:-2]).bold=True
        elif token.startswith('[') and '](' in token:
            label,url=token[1:-1].split('](',1); hyperlink(p,label,url)
        else:
            # Typeset symbolic indices while retaining ordinary prose and hyperlinks.
            pos=0
            for m in re.finditer(r'([A-Za-zΑ-Ωα-ω])_([A-Za-zÀ-ÿ]+)',token):
                p.add_run(token[pos:m.start()]); p.add_run(m.group(1))
                p.add_run(m.group(2)).font.subscript=True
                pos=m.end()
            p.add_run(token[pos:])

def math_run(text):
    r=OxmlElement('m:r'); pr=OxmlElement('m:rPr'); pr.append(OxmlElement('m:nor')); r.append(pr)
    t=OxmlElement('m:t'); t.set(qn('xml:space'),'preserve'); t.text=text; r.append(t)
    return r

def math_content(math,text):
    pos=0
    for m in re.finditer(r'([A-Za-zΑ-Ωα-ω])_([A-Za-zÀ-ÿ]+)',text):
        if pos<m.start(): math.append(math_run(text[pos:m.start()]))
        sub=OxmlElement('m:sSub'); e=OxmlElement('m:e'); e.append(math_run(m.group(1)))
        idx=OxmlElement('m:sub'); idx.append(math_run(m.group(2))); sub.append(e); sub.append(idx); math.append(sub)
        pos=m.end()
    if pos<len(text): math.append(math_run(text[pos:]))

def table(rows):
    n=len(rows[0]); t=doc.add_table(rows=1,cols=n); t.autofit=False
    # All tables fit the 6.9 inch text column.
    if n==2: widths=[2.35,4.55]
    else: widths=[1.7,2.65,2.55]
    if rows[0][0]=='Illustrative quantity': widths=[3.1,1.9,1.9]
    if rows[0][0]=='Symbol': widths=[1.35,5.55]
    if rows[0][0]=='00M component': widths=[1.1,2.7,3.1]
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    props.append(borders)
    for idx,values in enumerate(rows):
        row=t.rows[0] if idx==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx==0: trpr.append(OxmlElement('w:tblHeader'))
        for j,(c,value) in enumerate(zip(row.cells,values)):
            c.width=Inches(widths[j]); pr=c._tc.get_or_add_tcPr()
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            if idx==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E2ECF2'); pr.append(shade)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.04
            inline(p,value)
            for run in p.runs:
                run.font.size=Pt(10)
                if idx==0: run.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=[line for line in src.read_text().splitlines() if not re.fullmatch(r'<a id="[^"]+"></a>',line)]; i=0; tables=0; equations=0; part=''
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in cells): rows.append(cells)
            i+=1
        table(rows); tables+=1; continue
    if line.startswith('# '):
        part=line[2:].split(' ')[0]
        if i==0: doc.add_paragraph(line[2:],'Title')
        else: doc.add_paragraph(line[2:],'Heading 1')
    elif line.startswith('## '): doc.add_paragraph(line[3:],'Heading 2')
    elif line.startswith('!['):
        match=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
        if not match: raise ValueError(line)
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(4)
        pic=p.add_run().add_picture(str(ROOT/match.group(2)),width=Inches(6.9))
        pic._inline.docPr.set('descr',match.group(1))
    elif line.startswith('Figure '):
        p=doc.add_paragraph(); inline(p,line)
        p.paragraph_format.line_spacing=1.05; p.paragraph_format.space_after=Pt(9)
        for run in p.runs: run.font.size=Pt(9)
    elif line.startswith('> '):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18)
        math=OxmlElement('m:oMath'); math_content(math,line[2:]); p._p.append(math); equations+=1
    elif i==2: doc.add_paragraph(line,'Subtitle')
    else:
        p=doc.add_paragraph()
        if part=='1':
            p.paragraph_format.line_spacing=1.08
            p.paragraph_format.space_after=Pt(4)
        if re.match(r'^[1-9]\. ',line):
            p.paragraph_format.left_indent=Inches(.22)
            p.paragraph_format.first_line_indent=Inches(-.22)
        inline(p,line)
        if i in (4,6):
            for r in p.runs: r.font.size=Pt(10)
    i+=1
out=ROOT/'Escenario-creatividad-validacion.docx'; doc.save(out)
print({'output':str(out),'words':len(src.read_text().split()),'tables':tables,'equations':equations})
