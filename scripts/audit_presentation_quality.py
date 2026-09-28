#!/usr/bin/env python3
"""Audit the two canonical Ecosystem Positioning v1.2 PowerPoint decks.

The audit is deliberately independent of presentation claims. It checks that the
actual PPTX packages committed to GitHub are structurally readable, can be opened
and rendered by a desktop office engine, keep all shapes within the slide canvas,
and expose potential text-density/overlap risks for human visual review.

Generated outputs are written to ``presentation-audit/`` for CI artifact upload.
"""
from __future__ import annotations

import json
import math
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
PRESENTATION_DIR = ROOT / "presentations" / "ecosystem-positioning"
OUTPUT = ROOT / "presentation-audit"

DECKS = [
    PRESENTATION_DIR / "Ecosystem_Positioning_Requirements_Evidence_Canonical_v1.2.pptx",
    PRESENTATION_DIR / "Ecosystem_Positioning_Architecture_Implementation_Canonical_v1.2.pptx",
]

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

EMU_PER_INCH = 914400
PT_PER_INCH = 72


def run(cmd: list[str], *, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        cmd,
        cwd=str(cwd) if cwd else None,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )


def rel_target(base: str, target: str) -> str:
    """Resolve an OOXML relationship target relative to a package part."""
    from posixpath import dirname, normpath, join
    return normpath(join(dirname(base), target))


def validate_zip_and_relationships(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {"zip_ok": False, "xml_ok": False, "relationship_errors": []}
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        if bad:
            result["zip_bad_member"] = bad
            return result
        result["zip_ok"] = True
        names = set(zf.namelist())
        required = {
            "[Content_Types].xml",
            "_rels/.rels",
            "ppt/presentation.xml",
            "ppt/_rels/presentation.xml.rels",
        }
        missing = sorted(required - names)
        if missing:
            result["missing_required_parts"] = missing
            return result

        # Parse every XML/RELS part. A malformed package part is a hard failure.
        parse_errors: list[str] = []
        for name in sorted(names):
            if name.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(zf.read(name))
                except Exception as exc:  # pragma: no cover - diagnostic path
                    parse_errors.append(f"{name}: {exc}")
        result["xml_parse_errors"] = parse_errors
        if parse_errors:
            return result

        # Verify internal relationship targets exist. External links are excluded.
        rel_errors: list[str] = []
        for rel_name in sorted(n for n in names if n.endswith(".rels")):
            root = ET.fromstring(zf.read(rel_name))
            if rel_name == "_rels/.rels":
                source_part = ""
            else:
                # ppt/slides/_rels/slide1.xml.rels -> ppt/slides/slide1.xml
                prefix, filename = rel_name.rsplit("/_rels/", 1)
                source_part = f"{prefix}/{filename[:-5]}"  # remove .rels
            for rel in root.findall(f"{{{REL_NS}}}Relationship"):
                if rel.attrib.get("TargetMode") == "External":
                    continue
                target = rel.attrib.get("Target", "")
                resolved = rel_target(source_part, target) if source_part else target.lstrip("/")
                resolved = resolved.lstrip("/")
                if resolved not in names:
                    rel_errors.append(f"{rel_name}: {rel.attrib.get('Id')} -> {target} ({resolved})")
        result["relationship_errors"] = rel_errors
        result["xml_ok"] = not rel_errors

        slides = sorted(
            (n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)", n).group(1)),
        )
        result["slide_parts"] = len(slides)
    return result


def shape_text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def effective_font_points(shape: Any) -> float:
    sizes: list[float] = []
    if getattr(shape, "has_text_frame", False):
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    sizes.append(r.font.size.pt)
    return min(sizes) if sizes else 18.0


def intersection_area(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> int:
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    w = max(0, min(ax2, bx2) - max(ax1, bx1))
    h = max(0, min(ay2, by2) - max(ay1, by1))
    return w * h


def analyze_layout(path: Path) -> dict[str, Any]:
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    slides: list[dict[str, Any]] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        outside: list[dict[str, Any]] = []
        low_font: list[dict[str, Any]] = []
        dense_text: list[dict[str, Any]] = []
        text_shapes: list[dict[str, Any]] = []

        for shape_idx, shape in enumerate(slide.shapes, start=1):
            x, y, w, h = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
            right, bottom = x + w, y + h
            text = shape_text(shape)
            record = {
                "shape": shape_idx,
                "name": getattr(shape, "name", ""),
                "x": x,
                "y": y,
                "w": w,
                "h": h,
                "text": text[:240],
            }
            # Small tolerance for shadows/bleeds and intentionally full-bleed artwork.
            tol = int(0.03 * EMU_PER_INCH)
            if x < -tol or y < -tol or right > sw + tol or bottom > sh + tol:
                outside.append(record)

            if text:
                pt = effective_font_points(shape)
                record["min_font_pt"] = pt
                if pt < 10:
                    low_font.append(record)

                # Conservative density heuristic. It does not claim overflow; it
                # marks candidates for rendered visual review.
                width_in = max(w / EMU_PER_INCH, 0.1)
                height_in = max(h / EMU_PER_INCH, 0.1)
                chars = len(text)
                line_breaks = text.count("\n") + 1
                estimated_chars_per_line = max(8.0, width_in * 72.0 / max(pt * 0.52, 4.0))
                estimated_lines = max(line_breaks, math.ceil(chars / estimated_chars_per_line))
                line_height_in = max(pt * 1.18 / PT_PER_INCH, 0.12)
                estimated_height = estimated_lines * line_height_in
                record["estimated_text_height_in"] = round(estimated_height, 3)
                record["box_height_in"] = round(height_in, 3)
                if estimated_height > height_in * 1.08:
                    dense_text.append(record)
                text_shapes.append(record)

        overlaps: list[dict[str, Any]] = []
        # Only flag substantial overlap between *different text-bearing boxes*.
        # Background panels and decorative non-text shapes are ignored.
        for i in range(len(text_shapes)):
            a = text_shapes[i]
            abox = (a["x"], a["y"], a["x"] + a["w"], a["y"] + a["h"])
            for j in range(i + 1, len(text_shapes)):
                b = text_shapes[j]
                bbox = (b["x"], b["y"], b["x"] + b["w"], b["y"] + b["h"])
                area = intersection_area(abox, bbox)
                if area <= 0:
                    continue
                smaller = max(1, min(a["w"] * a["h"], b["w"] * b["h"]))
                ratio = area / smaller
                if ratio > 0.12:
                    overlaps.append({
                        "shape_a": a["shape"],
                        "shape_b": b["shape"],
                        "overlap_ratio_of_smaller": round(ratio, 3),
                        "text_a": a["text"][:120],
                        "text_b": b["text"][:120],
                    })

        slides.append({
            "slide": slide_idx,
            "outside_canvas": outside,
            "font_below_10pt": low_font,
            "text_density_candidates": dense_text,
            "text_box_overlap_candidates": overlaps,
            "shape_count": len(slide.shapes),
        })

    return {
        "slide_count": len(prs.slides),
        "slide_width": sw,
        "slide_height": sh,
        "slides": slides,
    }


def libreoffice_render(path: Path, deck_out: Path) -> dict[str, Any]:
    deck_out.mkdir(parents=True, exist_ok=True)
    pdf_dir = deck_out / "pdf"
    png_dir = deck_out / "png"
    pdf_dir.mkdir(exist_ok=True)
    png_dir.mkdir(exist_ok=True)

    proc = run([
        "libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir), str(path)
    ])
    pdf = pdf_dir / f"{path.stem}.pdf"
    result: dict[str, Any] = {
        "libreoffice_returncode": proc.returncode,
        "libreoffice_output": proc.stdout[-4000:],
        "pdf_created": pdf.exists(),
    }
    if proc.returncode != 0 or not pdf.exists():
        return result

    info = run(["pdfinfo", str(pdf)])
    result["pdfinfo"] = info.stdout[-4000:]
    page_match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.M)
    result["pdf_pages"] = int(page_match.group(1)) if page_match else None

    render = run([
        "pdftoppm", "-png", "-r", "144", str(pdf), str(png_dir / "slide")
    ])
    result["pdftoppm_returncode"] = render.returncode
    result["pdftoppm_output"] = render.stdout[-4000:]
    images = sorted(png_dir.glob("slide-*.png"))
    result["rendered_pngs"] = len(images)
    if images:
        build_contact_sheet(images, deck_out / "contact-sheet.png")
    return result


def build_contact_sheet(images: list[Path], output: Path) -> None:
    thumbs: list[Image.Image] = []
    labels: list[str] = []
    max_w = 520
    for idx, path in enumerate(images, start=1):
        im = Image.open(path).convert("RGB")
        ratio = max_w / im.width
        thumb = im.resize((max_w, max(1, int(im.height * ratio))))
        thumbs.append(thumb)
        labels.append(str(idx))
    cols = 2
    pad = 22
    label_h = 34
    cell_w = max_w + 2 * pad
    cell_h = max(im.height for im in thumbs) + label_h + 2 * pad
    rows = math.ceil(len(thumbs) / cols)
    sheet = Image.new("RGB", (cols * cell_w, rows * cell_h), "white")
    draw = ImageDraw.Draw(sheet)
    for i, im in enumerate(thumbs):
        col, row = i % cols, i // cols
        x = col * cell_w + pad
        y = row * cell_h + pad + label_h
        draw.text((x, row * cell_h + pad), f"Slide {labels[i]}", fill="black")
        sheet.paste(im, (x, y))
    sheet.save(output)


def main() -> int:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    summary: dict[str, Any] = {"decks": {}}
    hard_errors: list[str] = []

    for deck in DECKS:
        if not deck.exists():
            hard_errors.append(f"missing canonical deck: {deck.relative_to(ROOT)}")
            continue
        deck_out = OUTPUT / deck.stem
        deck_out.mkdir(parents=True, exist_ok=True)
        shutil.copy2(deck, deck_out / deck.name)

        structural = validate_zip_and_relationships(deck)
        try:
            layout = analyze_layout(deck)
            python_pptx_open = True
            python_pptx_error = None
        except Exception as exc:
            layout = {}
            python_pptx_open = False
            python_pptx_error = repr(exc)
        rendered = libreoffice_render(deck, deck_out)

        data = {
            "file": str(deck.relative_to(ROOT)),
            "bytes": deck.stat().st_size,
            "structural": structural,
            "python_pptx_open": python_pptx_open,
            "python_pptx_error": python_pptx_error,
            "layout": layout,
            "render": rendered,
        }
        summary["decks"][deck.name] = data
        (deck_out / "audit.json").write_text(json.dumps(data, indent=2), encoding="utf-8")

        if not structural.get("zip_ok") or not structural.get("xml_ok"):
            hard_errors.append(f"OOXML structural failure: {deck.name}")
        if not python_pptx_open:
            hard_errors.append(f"python-pptx cannot open: {deck.name}: {python_pptx_error}")
        expected = layout.get("slide_count")
        pages = rendered.get("pdf_pages")
        pngs = rendered.get("rendered_pngs")
        if rendered.get("libreoffice_returncode") != 0 or not rendered.get("pdf_created"):
            hard_errors.append(f"LibreOffice cannot render: {deck.name}")
        elif expected is not None and pages != expected:
            hard_errors.append(f"PDF page mismatch: {deck.name}: {pages} != {expected}")
        elif expected is not None and pngs != expected:
            hard_errors.append(f"PNG render mismatch: {deck.name}: {pngs} != {expected}")

    summary["hard_errors"] = hard_errors
    (OUTPUT / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(json.dumps(summary, indent=2))
    if hard_errors:
        print("PRESENTATION QUALITY AUDIT: HARD FAIL")
        return 1
    print("PRESENTATION QUALITY AUDIT: STRUCTURAL/RENDER PASS; visual warnings require review of artifact renders")
    return 0


if __name__ == "__main__":
    sys.exit(main())
